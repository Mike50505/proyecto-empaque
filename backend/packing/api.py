import json
import uuid
from datetime import date
from io import BytesIO
from django.contrib.auth import authenticate, login, logout
from django.db import transaction
from django.db.models import Sum
from django.http import HttpResponse
from django.middleware.csrf import get_token
from django.utils import timezone
from django.core.exceptions import ObjectDoesNotExist
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework.exceptions import ValidationError
from django.views.decorators.csrf import csrf_protect
from django.views.decorators.cache import never_cache
from django.conf import settings
from . import services as s, importing, universe, program_deletion
from .models import *
from .part_validation import validate_part_customer, validate_program
from .client_groups import CLIENT_GROUPS, client_group, client_filter

def guarded(fn):
    from functools import wraps
    @wraps(fn)
    def wrap(*args,**kwargs):
        try: return fn(*args,**kwargs)
        except ObjectDoesNotExist: raise ValidationError('No se encontró el registro solicitado o la nómina activa.')
        except (ValueError,TypeError,KeyError) as e: raise ValidationError('Datos inválidos o incompletos; verifica el formulario.') from e
    return wrap

def roles(user): return {r:user.is_superuser or user.has_perm('packing.'+r) for r in ['consult','capture','plan','release']}

@api_view(['GET','POST','DELETE'])
@permission_classes([AllowAny])
@csrf_protect
def session(request):
    if request.method=='POST':
        user=authenticate(request,username=request.data.get('username'),password=request.data.get('password'))
        if not user: return Response({'detail':'Usuario o contraseña incorrectos.'},status=400)
        login(request,user)
    elif request.method=='DELETE': logout(request)
    return Response({'authenticated':request.user.is_authenticated,'username':request.user.get_username(),'roles':roles(request.user),'csrf':get_token(request),'release_mode':settings.RELEASE_MODE})

@api_view(['GET'])
def catalog(request):
    s.require(request.user,'consult')
    state=IntegrationState.objects.filter(name='universo').first()
    programs=Program.objects.filter(**({} if roles(request.user)['plan'] else {'published':True}))
    parts=[{**part,'universe_customer':part['customer'],'customer':client_group(part['customer'])} for part in Part.objects.order_by('number').values()]
    program_rows=[{**program,'client':client_group(program['client'])} for program in programs.values()]
    return Response({'clients':list(CLIENT_GROUPS),'part_clients':list(CLIENT_GROUPS),'current_year':timezone.localdate().year,'groups':list(BalanceGroup.objects.values()),'employees':list(Employee.objects.filter(active=True).values()),'shifts':list(Shift.objects.values()),'parts':parts,'programs':program_rows,'integration':{'last_sync':state.last_sync if state else None,'error':state.error if state else '', 'local_parts':Part.objects.count()},'today':str(timezone.localdate())})

@api_view(['POST'])
@guarded
@transaction.atomic
def employee_create(request):
    s.require(request.user,'plan');s.lock_ledger()
    payroll=str(request.data.get('payroll','')).strip()
    name=str(request.data.get('name','')).strip()
    if not payroll or len(payroll)>40 or not name or len(name)>150:
        raise ValidationError('Ingresa una nómina de hasta 40 caracteres y el nombre del trabajador (hasta 150).')
    existing=Employee.objects.filter(payroll=payroll).first()
    if existing:
        if existing.name!=name or not existing.active:raise ValidationError('La nómina ya existe con otros datos; revísala en Administración.')
        return Response({'id':existing.pk,'payroll':existing.payroll,'name':existing.name})
    employee=Employee.objects.create(payroll=payroll,name=name)
    s.audit(request.user,'employee_create',employee,payroll=payroll,name=name)
    return Response({'id':employee.pk,'payroll':employee.payroll,'name':employee.name},status=201)

@api_view(['GET'])
@guarded
def balance(request): return Response(s.balance(request.user,request.query_params))

def event_data(e):
    return {'id':e.pk,'key':str(e.key),'quantity':e.quantity,'part':e.target.attributes_snapshot.get('number',e.target.part.number),'demand':e.target_id,'client':client_group(e.target.program.client),'scope':e.target.program.scope,'year':e.target.program.year,'week':e.target.program.week,'effective_at':e.effective_at,'created_at':e.created_at,'operational_date':e.operational_date,'shift':e.shift.code,'packer':e.packer.payroll,'inspector':e.inspector.payroll if e.inspector else None,'actor':e.actor.get_username(),'released':e.released,'release_mode':e.release_mode,'voided':e.voided,'credit':s.available(e),'allocations':list(e.allocations.values('demand','quantity','active','created_at')),'notes':e.notes}

@api_view(['GET','POST'])
@guarded
def captures(request):
    if request.method=='POST': return Response(event_data(s.capture(request.user,request.data)),status=201)
    s.require(request.user,'consult')
    qs=Capture.objects.select_related('target__part','target__program','packer','inspector','shift','actor').order_by('-id')
    allowed=s.demand_queryset(request.user,{})
    qs=qs.filter(target__in=allowed)
    if request.query_params.get('date'): qs=qs.filter(operational_date=request.query_params['date'])
    if request.query_params.get('shift'): qs=qs.filter(shift__code=request.query_params['shift'])
    if request.query_params.get('packer'): qs=qs.filter(packer__payroll=request.query_params['packer'])
    if request.query_params.get('inspector'): qs=qs.filter(inspector__payroll=request.query_params['inspector'])
    if request.query_params.get('actor'): qs=qs.filter(actor__username=request.query_params['actor'])
    if request.query_params.get('state')=='pending':qs=qs.filter(released=False,voided=False)
    if request.query_params.get('state')=='released':qs=qs.filter(released=True,voided=False)
    if request.query_params.get('state')=='voided':qs=qs.filter(voided=True)
    if request.query_params.get('search'): qs=qs.filter(target__part__number__icontains=request.query_params['search'])
    if request.query_params.get('week'): qs=qs.filter(target__program__week=request.query_params['week'])
    if request.query_params.get('year'): qs=qs.filter(target__program__year=request.query_params['year'])
    if request.query_params.get('client'): qs=qs.filter(client_filter('target__program__client',request.query_params['client']))
    if request.query_params.get('subclient'): qs=qs.filter(target__part__customer__iexact=request.query_params['subclient'])
    if request.query_params.get('scope'): qs=qs.filter(target__program__scope=request.query_params['scope'])
    if request.query_params.get('group'): qs=qs.filter(target__program__group_id=request.query_params['group'])
    return Response([event_data(e) for e in qs])

@api_view(['GET'])
@guarded
def analytics(request):
    ids=list(s.demand_queryset(request.user,request.query_params).values_list('id',flat=True))
    qs=Capture.objects.filter(target_id__in=ids,voided=False)
    if request.query_params.get('date'):qs=qs.filter(operational_date=request.query_params['date'])
    if request.query_params.get('shift'):qs=qs.filter(shift__code=request.query_params['shift'])
    physical=list(qs.values('operational_date','shift__code').annotate(quantity=Sum('quantity')).order_by('operational_date','shift__code'))
    trend=[]
    for snap in Snapshot.objects.order_by('-cutoff')[:60]:
        selected=[r for r in snap.payload['rows'] if r['id'] in ids]
        if not selected: continue
        quantity=sum(r['quantity'] for r in selected);done=sum(r['applied'] for r in selected)
        trend.append({'id':snap.pk,'cutoff':snap.cutoff,'version':snap.version,'quantity':quantity,'applied':done,'progress':100*done/quantity if quantity else None})
    return Response({'physical':physical,'pending_release':qs.filter(released=False).aggregate(n=Sum('quantity'))['n'] or 0,'trend':list(reversed(trend)),'revisions':list(Revision.objects.filter(demand_id__in=ids).values('demand','before','after','reason','created_at').order_by('-id')[:100])})

@api_view(['POST'])
@guarded
def capture_action(request,pk,action):
    if action=='release': return Response(event_data(s.release(request.user,pk,request.data['inspector'])))
    if action=='correct':
        e=s.correct(request.user,pk,request.data['quantity'],request.data['reason'],request.data['key'])
        return Response(event_data(e) if e else {'voided':True})
    if action=='request-correction':
        s.require(request.user,'capture')
        event=Capture.objects.get(pk=pk)
        if not event.target.program.published: raise ValidationError('Programa no disponible en piso.')
        reason=request.data.get('reason','').strip()
        if not reason: raise ValidationError('Indica el motivo.')
        obj=CorrectionRequest.objects.create(capture=event,reason=reason,actor=request.user)
        return Response({'id':obj.pk})
    raise ValidationError('Acción desconocida.')

@api_view(['POST'])
@guarded
@transaction.atomic
def programs(request):
    s.require(request.user,'plan'); s.lock_ledger(); s.check_open(timezone.localdate())
    data=request.data; year=timezone.localdate().year
    client=str(data.get('client','')).strip()
    if not client or len(client)>150: raise ValidationError('Indica el cliente del programa.')
    if not data.get('lines'): raise ValidationError('Agrega al menos un renglón.')
    parts=[]
    for item in data['lines']:
        part=Part.objects.get(pk=item['part'])
        validate_part_customer(part,client)
        parts.append(part)
    client=client_group(parts[0].customer)
    start=date.fromisocalendar(year,int(data['week']),1)
    group,_=BalanceGroup.objects.get_or_create(name=importing.client_group_name(client))
    if not group.active: raise ValidationError('Cliente desactivado para nuevos programas.')
    p=Program.objects.create(group=group,client=client,scope=data.get('scope',''),year=year,week=int(data['week']),starts_on=start,ends_on=date.fromisocalendar(year,int(data['week']),7),reference='PROG-'+uuid.uuid4().hex.upper())
    if not data.get('lines'): raise ValidationError('Agrega al menos un renglón.')
    refs=set()
    for item,part in zip(data['lines'],parts):
        ref='ORD-'+uuid.uuid4().hex[:16].upper()
        validate_part_customer(part,client)
        if ref in refs: raise ValidationError('Referencia de renglón duplicada.')
        refs.add(ref)
        n=s.positive(item['quantity'],True)
        line=Demand.objects.create(program=p,part=part,reference=ref,original=n,quantity=n,priority=item.get('priority',2),due_on=item.get('due_on') or None,attributes_snapshot={'diameter':part.diameter,'attributes':part.attributes,'number':part.number})
        line.reference=f'ORD-{year}-{line.pk:06d}'
        line.save(update_fields=['reference'])
    s.audit(request.user,'create_program',p)
    return Response({'id':p.pk,'published':False},status=201)

@api_view(['POST'])
@guarded
@transaction.atomic
def program_action(request,pk):
    s.require(request.user,'plan'); s.lock_ledger()
    p=Program.objects.select_for_update().get(pk=pk)
    if request.data.get('published'): validate_program(p)
    if request.data.get('published') and not p.group.active: raise ValidationError('Grupo desactivado.')
    p.published=bool(request.data['published']); p.save(update_fields=['published'])
    s.audit(request.user,'publication',p,published=p.published)
    return Response({'id':p.pk,'published':p.published})


@api_view(['GET','DELETE'])
@guarded
def program_detail(request,pk):
    if request.method=='GET': return Response(program_deletion.inspect(request.user,pk))
    return Response(program_deletion.delete(request.user,pk,request.data.get('reason'),request.data.get('confirmation')))

@api_view(['POST'])
@guarded
@transaction.atomic
def revise(request,pk):
    line=s.revise(request.user,pk,request.data['quantity'],request.data['reason'])
    before={'priority':line.priority,'due_on':str(line.due_on) if line.due_on else None}
    if 'priority' in request.data:
        priority=s.positive(request.data['priority'])
        if priority not in [1,2,3]:raise ValidationError('Prioridad inválida.')
        line.priority=priority
    if 'due_on' in request.data:line.due_on=date.fromisoformat(request.data['due_on']) if request.data['due_on'] else None
    line.save(update_fields=['priority','due_on'])
    s.audit(request.user,'priority_due',line,request.data['reason'],before=before,after={'priority':line.priority,'due_on':str(line.due_on) if line.due_on else None})
    return Response({'id':line.pk,'quantity':line.quantity})

@api_view(['POST'])
@guarded
def transfer(request): return Response({'allocated':s.transfer(request.user,request.data['source'],request.data['demand'],request.data['quantity'],request.data['reason'],request.data.get('opening',False))})

@api_view(['POST'])
@guarded
def apply_surplus(request):
    from .surplus import apply_all
    return Response(apply_all(request.user))

@api_view(['GET','POST'])
@guarded
def imports(request):
    s.require(request.user,'plan')
    if request.method=='GET': return Response(list(ImportBatch.objects.order_by('-id').values('id','digest','mode','mapping','applied','created_at')[:100]))
    batch=importing.preview(request.user,request.FILES['file'],request.data['mode'],json.loads(request.data.get('mapping','{}')))
    return Response(importing.result(batch))

@api_view(['GET','POST'])
@guarded
def import_detail(request,pk):
    s.require(request.user,'plan')
    batch=importing.commit(request.user,pk,publish=request.data.get('publish') is True) if request.method=='POST' else ImportBatch.objects.get(pk=pk)
    return Response(importing.result(batch))

@api_view(['GET','POST'])
@guarded
def cuts(request):
    s.require(request.user,'plan')
    if request.method=='GET': return Response(list(Snapshot.objects.order_by('-id').values('id','date','version','scope','cutoff')[:200]))
    snap=s.snapshot(request.user,request.data.get('scope',{}),request.data.get('close',False),request.data.get('reason',''))
    return Response({'id':snap.pk,'version':snap.version,'payload':snap.payload})

@api_view(['GET'])
@guarded
def cut_detail(request,pk):
    s.require(request.user,'plan')
    obj=Snapshot.objects.get(pk=pk)
    return Response({'id':obj.pk,'date':obj.date,'version':obj.version,'cutoff':obj.cutoff,'payload':obj.payload})

@api_view(['POST'])
@guarded
@transaction.atomic
def reopen(request):
    s.require(request.user,'plan'); s.lock_ledger()
    if not request.data.get('reason','').strip(): raise ValidationError('Indica motivo de reapertura.')
    obj=ClosedPeriod.objects.get(date=request.data['date'])
    s.audit(request.user,'reopen',obj,request.data['reason'],date=str(obj.date)); obj.delete()
    return Response({'reopened':True})

@api_view(['GET'])
def audit(request):
    s.require(request.user,'plan')
    return Response({'audit':list(Audit.objects.order_by('-id').values('id','kind','object_id','actor__username','reason','data','created_at')[:500]),'requests':list(CorrectionRequest.objects.filter(resolved=False).values()),'closed':list(ClosedPeriod.objects.values())})

@api_view(['POST'])
def sync(request):
    s.require(request.user,'plan'); return Response(universe.sync())

def safe(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')): return "'"+value
    return value

def xlsx_response(rows,filename):
    w=Workbook(); ws=w.active; ws.title='Empaques MESA'
    for row in rows: ws.append([safe(v) for v in row])
    for c in ws[1]: c.font=Font(color='FFFFFF',bold=True); c.fill=PatternFill('solid',fgColor='163D33')
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(45,max(16,max(len(str(c.value or '')) for c in col)+2))
    out=BytesIO(); w.save(out)
    response=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition']=f'attachment; filename="{filename}.xlsx"'; return response

@api_view(['GET'])
@guarded
def export(request):
    data=s.balance(request.user,request.query_params)
    if request.query_params.get('snapshot'):
        s.require(request.user,'plan'); data=Snapshot.objects.get(pk=request.query_params['snapshot']).payload
    if 'order_state' in request.query_params and not request.query_params.get('snapshot'):
        keys=['client','scope','year','week','reference','part','diameter','quantity','packed','pending_release','applied','packing_pending','pending','credit']
        headers=['Cliente','Destino','Año','Semana','Folio de orden','Número de parte','Diámetro','Programado','Empacado','Por liberar','Aplicado','Por empacar','Pendiente por cumplir','Excedente']
        return xlsx_response([headers]+[[r[k] for k in keys] for r in data['rows']],'ordenes-empaque')
    keys=['group','client','scope','year','week','reference','part','diameter','original','quantity','applied','pending','credit','A','B']
    return xlsx_response([['Grupo','Cliente','Ámbito','Año','Semana origen','Renglón','Número de parte','Diámetro','Original','Vigente','Aplicado','Pendiente','Excedente','A del día','B del día']]+[[r[k] for k in keys] for r in data['rows']],'balance-empaques')

@api_view(['GET'])
@never_cache
def template(request):
    s.require(request.user,'plan')
    columns=['cliente','destino_o_ambito','semana_programa','numero_parte','cantidad_programada']
    if request.query_params.get('mode')=='revision': columns.append('referencia')
    filename='plantilla-revision-ordenes' if request.query_params.get('mode')=='revision' else 'plantilla-ordenes-por-cliente'
    return xlsx_response([columns],filename)

@api_view(['GET'])
@guarded
def production_export(request):
    s.require(request.user,'consult')
    qs=Capture.objects.filter(target__in=s.demand_queryset(request.user,request.query_params)).select_related('target__part','packer','inspector','actor','shift')
    if request.query_params.get('date'): qs=qs.filter(operational_date=request.query_params['date'])
    if request.query_params.get('shift'): qs=qs.filter(shift__code=request.query_params['shift'])
    return xlsx_response([['Folio','Pieza','Cantidad','Fecha efectiva','Registro','Fecha operativa','Turno','Empacador','Inspector','Capturista','Liberada','Anulada']]+[[e.pk,e.target.part.number,e.quantity,e.effective_at.isoformat(),e.created_at.isoformat(),str(e.operational_date),e.shift.code,e.packer.payroll,e.inspector.payroll if e.inspector else '',e.actor.username,e.released,e.voided] for e in qs],'movimientos-empaque')
