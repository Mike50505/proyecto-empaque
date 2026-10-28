"""Aceptación con datos sintéticos y libros reales intactos. PostgreSQL obligatorio."""
import json
import uuid
import time
import tempfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date,timedelta,datetime,time as dt_time
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from django.test import TransactionTestCase, override_settings
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.db import connection,close_old_connections
from django.test.utils import CaptureQueriesContext
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook,load_workbook
from rest_framework.test import APIClient
from rest_framework.exceptions import ValidationError,PermissionDenied
from .models import *
from . import services as s, importing

class BusinessAcceptance(TransactionTestCase):
    reset_sequences=True
    @classmethod
    def setUpClass(cls):
        cls.temp_media=tempfile.TemporaryDirectory(prefix='empaques-tests-')
        cls.media_settings=override_settings(MEDIA_ROOT=cls.temp_media.name)
        cls.media_settings.enable()
        super().setUpClass()
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        cls.media_settings.disable();cls.temp_media.cleanup()
    def setUp(self):
        assert connection.vendor=='postgresql','Estas pruebas exigen PostgreSQL real.'
        self.user=get_user_model().objects.create_superuser(username='synthetic-planner',password='synthetic-only')
        self.group=BalanceGroup.objects.create(name='SINTÉTICO')
        self.part=Part.objects.create(number='0001-X',customer='CLIENTE-SINTÉTICO',universe_id=str(uuid.uuid4()),provisional=False,attributes={'active':True})
        self.shift=Shift.objects.create(code='A')
        Shift.objects.create(code='B')
        Employee.objects.create(payroll='TEST-P',name='Empacador sintético')
        Employee.objects.create(payroll='TEST-I',name='Inspector sintético',inspector=True)
        self.line=self.make_line(300)
        self.client=APIClient(); self.client.force_authenticate(self.user)

    def make_line(self,quantity,week=42,client='CLIENTE-SINTÉTICO',scope='SLP',part=None):
        p=Program.objects.create(group=self.group,client=client,scope=scope,year=2026,week=week,starts_on=date.fromisocalendar(2026,week,1),ends_on=date.fromisocalendar(2026,week,7),reference=str(uuid.uuid4()),published=True)
        return Demand.objects.create(program=p,part=part or self.part,reference='1',original=quantity,quantity=quantity)

    def payload(self,quantity,line=None,shift='A',key=None):
        workday=timezone.localdate()-timedelta(days=1)
        effective=timezone.make_aware(datetime.combine(workday,dt_time(10 if shift=='A' else 18)),timezone.get_current_timezone())
        return {'key':str(key or uuid.uuid4()),'demand':(line or self.line).pk,'quantity':quantity,'effective_at':effective.isoformat(),'operational_date':str(timezone.localdate()),'shift':shift,'packer':'TEST-P','inspector':'TEST-I','accept_excess':True}

    def pack(self,quantity,line=None,shift='A'): return s.capture(self.user,self.payload(quantity,line,shift))

    def pending_capture(self,quantity):
        # Historical rows predate release in the capture form.
        event=self.pack(quantity)
        event.allocations.all().delete()
        event.released=False;event.released_by=None;event.inspector_payroll=''
        event.save()
        return event


    def test_01_partial_close_persisted(self):
        self.pack(200);self.assertEqual(s.applied(self.line),200)
        self.pack(100,shift='B');self.assertEqual(s.applied(Demand.objects.get(pk=self.line.pk)),300)
        self.assertEqual(s.balance(self.user,{})['totals']['progress'],100)

    def test_02_incremental(self):
        self.pack(200);self.pack(50);self.assertEqual(s.applied(self.line),250)

    @override_settings(RELEASE_MODE='declared')
    def test_capture_accepts_unregistered_payrolls_and_releases_inline(self):
        payload=self.payload(10);payload['packer']='TEMP-UNREGISTERED-55';payload['inspector']='TEST-I'
        event=s.capture(self.user,payload)
        self.assertEqual(event.packer_payroll,'TEMP-UNREGISTERED-55')
        self.assertIsNone(event.packer)
        self.assertIsNone(event.inspector)
        self.assertTrue(event.released)
        self.assertEqual(event.inspector_payroll,'TEST-I')
        self.assertEqual(event.released_by,self.user)
        self.assertEqual(s.applied(self.line),10)

    def test_release_accepts_unregistered_inspector_payroll(self):
        event=self.pending_capture(10)
        released=s.release(self.user,event.pk,'QA-TEMP-998')
        self.assertTrue(released.released)
        self.assertEqual(released.inspector_payroll,'QA-TEMP-998')
        self.assertIsNone(released.inspector)

    def test_inline_release_requires_payroll_and_leaves_no_partial_capture(self):
        for value in ['', '   ', None, 'X'*41]:
            payload=self.payload(10);payload['inspector']=value
            with self.assertRaises(ValidationError):
                s.capture(self.user,payload)
        self.assertEqual(Capture.objects.count(),0)
        self.assertEqual(s.applied(self.line),0)

    @override_settings(RELEASE_MODE='confirmation')
    def test_capture_operator_declares_release_in_same_transaction(self):
        operator=get_user_model().objects.create_user('capture-only')
        operator.user_permissions.add(Permission.objects.get(codename='capture'))
        payload=self.payload(25);payload['inspector']='  00045  '
        event=s.capture(operator,payload)
        self.assertTrue(event.released)
        self.assertEqual(event.release_mode,'declared')
        self.assertEqual(event.inspector_payroll,'00045')
        self.assertEqual(event.released_by,operator)
        self.assertEqual(s.applied(self.line),25)
        self.assertEqual(s.capture(operator,payload).pk,event.pk)
        self.assertEqual(s.applied(self.line),25)

    def test_capture_api_returns_release_payroll_and_applied_quantity(self):
        payload=self.payload(20);payload['inspector']='QA-UNREGISTERED'
        response=self.client.post('/api/captures/',payload,format='json')
        self.assertEqual(response.status_code,201,response.data)
        self.assertTrue(response.data['released'])
        self.assertEqual(response.data['inspector'],'QA-UNREGISTERED')
        self.assertEqual(s.applied(self.line),20)

    def test_shift_is_inferred_from_effective_time_not_payload(self):
        payload=self.payload(10,shift='B');payload['shift']='A'
        event=s.capture(self.user,payload)
        self.assertEqual(event.shift.code,'B')

    def test_after_midnight_belongs_to_previous_night_shift_date(self):
        workday=timezone.localdate()-timedelta(days=1)
        effective=timezone.make_aware(datetime.combine(workday,dt_time(1,30)),timezone.get_current_timezone())
        payload=self.payload(10);payload['effective_at']=effective.isoformat()
        event=s.capture(self.user,payload)
        self.assertEqual(event.shift.code,'B')
        self.assertEqual(event.operational_date,workday-timedelta(days=1))

    def test_shift_schedule_boundaries_cover_every_hour(self):
        workday=timezone.localdate()-timedelta(days=1)
        cases=[(dt_time(0,0),'B'),(dt_time(2,15),'B'),(dt_time(2,16),'B'),(dt_time(4,37),'B'),(dt_time(4,38),'A'),(dt_time(6,59),'A'),(dt_time(7,0),'A'),(dt_time(16,36),'A'),(dt_time(16,37),'B'),(dt_time(23,59),'B')]
        for clock,expected in cases:
            payload=self.payload(1)
            payload['effective_at']=timezone.make_aware(datetime.combine(workday,clock),timezone.get_current_timezone()).isoformat()
            event=s.capture(self.user,payload)
            self.assertEqual(event.shift.code,expected,clock.isoformat())
            operational_day=workday-timedelta(days=1) if clock<=dt_time(4,37) else timezone.localdate()
            self.assertEqual(event.operational_date,operational_day,clock.isoformat())
    @override_settings(RELEASE_MODE='confirmation')
    def test_capture_sheet_counts_physical_before_release(self):
        event=self.pending_capture(200)
        row=s.balance(self.user,{})['rows'][0]
        self.assertEqual(row['packed'],200)
        self.assertEqual(row['applied'],0)
        self.assertEqual(row['quantity'],300)
        self.assertEqual(row['packing_pending'],100)
        self.assertEqual(row['pending_release'],200)
        self.assertEqual(row['pending'],300)
        s.correct(self.user,event.pk,150,'Cantidad real',uuid.uuid4())
        self.assertEqual(s.balance(self.user,{})['rows'][0]['packed'],150)
        self.assertEqual(s.balance(self.user,{})['rows'][0]['packing_pending'],150)
        self.assertEqual(s.balance(self.user,{})['rows'][0]['pending_release'],150)
        self.line.refresh_from_db();self.assertEqual(self.line.quantity,300)

    @override_settings(RELEASE_MODE='confirmation')
    def test_unified_orders_filters_export_and_release_totals(self):
        self.pending_capture(200)
        draft=self.make_line(90,43);draft.program.published=False;draft.program.save()
        data=s.balance(self.user,{'publication':'published','order_state':'release'})
        self.assertEqual([r['id'] for r in data['rows']],[self.line.pk])
        self.assertEqual(data['totals']['packed'],200)
        self.assertEqual(data['totals']['pending_release'],200)
        self.assertEqual(data['totals']['pending'],300)
        self.assertEqual(data['totals']['packing_pending'],100)
        response=self.client.get('/api/export/',{'publication':'published','order_state':'release'})
        sheet=load_workbook(BytesIO(response.content)).active
        records=list(sheet.values);self.assertEqual(len(records),2)
        exported=dict(zip(records[0],records[1]))
        self.assertEqual(exported['Empacado'],200)
        self.assertEqual(exported['Por liberar'],200)
        self.assertEqual(exported['Pendiente por cumplir'],300)
        self.assertEqual(s.balance(self.user,{'publication':'published','order_state':'complete'})['rows'],[])

    def test_unified_orders_excess_is_not_pending_release(self):
        self.pack(400)
        row=s.balance(self.user,{})['rows'][0]
        self.assertEqual(row['packed'],400);self.assertEqual(row['credit'],100)
        self.assertEqual(row['pending_release'],0);self.assertEqual(row['pending'],0)
        self.assertEqual(s.balance(self.user,{'order_state':'release'})['rows'],[])

    def test_floor_balance_preserves_transferred_credits(self):
        event=self.pack(400)
        future=self.make_line(200,43)
        s.transfer(self.user,event.pk,future.pk,100,'Aplicación de excedente')
        rows={r['id']:r for r in s.balance(self.user,{})['rows']}
        self.assertEqual(rows[self.line.pk]['quantity'],300)
        self.assertEqual(rows[self.line.pk]['packing_pending'],0)
        self.assertEqual(rows[future.pk]['quantity'],200)
        self.assertEqual(rows[future.pk]['packing_pending'],100)
        self.assertEqual(s.packing_pending(future),100)

    def test_03_excess_transfer_without_physical_duplicate(self):
        s.revise(self.user,self.line.pk,600,'Programa sintético')
        e=self.pack(700);self.assertEqual(s.available(e),100)
        future=self.make_line(400,43)
        s.transfer(self.user,e.pk,future.pk,100,'Aplicación sintética')
        self.assertEqual(s.applied(future),100);self.assertEqual(Capture.objects.count(),1)

    def test_04_different_parts_do_not_offset(self):
        self.pack(400)
        part=Part.objects.create(number='OTHER')
        self.make_line(100,part=part)
        self.assertEqual(s.balance(self.user,{})['totals']['pending'],100)

    def test_05_scope_and_client_compatibility(self):
        e=self.pack(400)
        for target in [self.make_line(100,scope='USA'),self.make_line(100,client='OTRO')]:
            with self.assertRaises(ValidationError):s.transfer(self.user,e.pk,target.pk,100,'No compatible')

    def test_06_future_commitment_different_from_physical_date(self):
        future=self.make_line(100,43);e=self.pack(100,future)
        self.assertEqual(e.operational_date,timezone.localdate());self.assertEqual(s.applied(self.line),0);self.assertEqual(s.applied(future),100)

    def test_07_carry_is_original_identity(self):
        old=self.make_line(200,41)
        rows=s.balance(self.user,{'year':'2026','week':'42','view':'operational'})['rows']
        self.assertEqual(sum(r['pending'] for r in rows),500);self.assertEqual(len({r['id'] for r in rows}),2)

    def template_file(self,qty=3200):
        w=Workbook();ws=w.active;ws.title='Programa'
        ws.append(['grupo','cliente','destino_o_ambito','anio_programa','semana_programa','numero_parte','cantidad_programada','referencia'])
        ws.append(['SINTÉTICO','CLIENTE-SINTÉTICO','SLP',2026,42,self.part.number,qty,'1'])
        out=BytesIO();w.save(out);return SimpleUploadedFile('synthetic.xlsx',out.getvalue())

    def test_08_import_idempotent(self):
        content=self.template_file().read()
        b=importing.preview(self.user,SimpleUploadedFile('synthetic.xlsx',content),'new',{'sheet':'Programa','reason':'Prueba'})
        importing.commit(self.user,b.pk)
        count=Demand.objects.count(); importing.commit(self.user,b.pk)
        b2=importing.preview(self.user,SimpleUploadedFile('synthetic.xlsx',content),'new',{'sheet':'Programa','reason':'Prueba'})
        self.assertEqual(b.pk,b2.pk);self.assertEqual(Demand.objects.count(),count)

    def simple_template(self,rows):
        workbook=Workbook();sheet=workbook.active;sheet.title='Programa'
        sheet.append(['cliente','destino_o_ambito','semana_programa','numero_parte','cantidad_programada'])
        for row in rows:sheet.append(row)
        stream=BytesIO();workbook.save(stream)
        return SimpleUploadedFile('simple.xlsx',stream.getvalue())

    def test_simple_template_merges_same_part_in_week_and_keeps_automatic_folios(self):
        file=self.simple_template([[self.part.customer,'SLP',41,self.part.number,100],[self.part.customer,'SLP',43,self.part.number,200],[self.part.customer,'SLP',43,self.part.number,300]])
        content=file.read()
        batch=importing.preview(self.user,SimpleUploadedFile('simple.xlsx',content),'new',{'reason':'Simple'})
        self.assertFalse([i for i in batch.preview['issues'] if i['severity']=='blocking'])
        self.assertEqual({r['year'] for r in batch.preview['rows']},{timezone.localdate().year})
        self.assertEqual({r['group'] for r in batch.preview['rows']},{self.part.customer})
        importing.commit(self.user,batch.pk,publish=True)
        batch.refresh_from_db()
        refs=[r['reference'] for r in batch.preview['rows']]
        self.assertEqual(len(set(refs)),2)
        for ref in refs:self.assertRegex(ref,r'^ORD-\d{4}-\d{6,}$')
        before=Demand.objects.count()
        self.assertEqual(before,3)
        self.assertEqual(Demand.objects.get(program__week=43).quantity,500)
        repeated=importing.preview(self.user,SimpleUploadedFile('simple.xlsx',content),'new',{'reason':'Simple'})
        importing.commit(self.user,repeated.pk,publish=True)
        self.assertEqual(repeated.pk,batch.pk);self.assertEqual(Demand.objects.count(),before)
        self.assertEqual(set(Demand.objects.filter(reference__in=refs).values_list('reference',flat=True)),set(refs))

    def test_template_has_only_operational_columns(self):
        response=self.client.get('/api/template/')
        sheet=load_workbook(BytesIO(response.content)).active
        self.assertEqual(list(next(sheet.values)),['numero_parte','cantidad_programada','semana_programa'])

    def test_new_import_can_sum_or_replace_same_part_in_existing_week(self):
        first=importing.preview(self.user,self.simple_template([[self.part.customer,'SLP',43,self.part.number,100]]),'new',{'reason':'First load'})
        importing.commit(self.user,first.pk)
        summed=importing.preview(self.user,self.simple_template([[self.part.customer,'SLP',43,self.part.number,60]]),'new',{'reason':'Add quantity','duplicate_action':'sum'})
        self.assertTrue(summed.preview['rows'][0]['duplicate_existing'])
        self.assertEqual(summed.preview['rows'][0]['quantity_after'],160)
        importing.commit(self.user,summed.pk)
        self.assertEqual(Demand.objects.get(program__week=43).quantity,160)
        replaced=importing.preview(self.user,self.simple_template([[self.part.customer,'SLP',43,self.part.number,30]]),'new',{'reason':'Replace quantity','duplicate_action':'replace'})
        self.assertEqual(replaced.preview['rows'][0]['quantity_after'],30)
        importing.commit(self.user,replaced.pk)
        self.assertEqual(Demand.objects.filter(program__week=43).count(),1)
        self.assertEqual(Demand.objects.get(program__week=43).quantity,30)

    def test_year_rollover_requires_a_new_preview(self):
        batch=importing.preview(self.user,self.simple_template([[self.part.customer,'SLP',43,self.part.number,100]]),'new',{'reason':'Year rollover'})
        count=Demand.objects.count()
        next_year=date(timezone.localdate().year+1,1,2)
        with patch('packing.importing.timezone.localdate',return_value=next_year):
            with self.assertRaises(ValidationError):importing.commit(self.user,batch.pk)
        self.assertEqual(Demand.objects.count(),count)

    def test_reupload_of_a_previously_applied_legacy_batch_does_not_duplicate_orders(self):
        file=self.simple_template([[self.part.customer,'SLP',43,self.part.number,100]])
        content=file.read()
        batch=importing.preview(self.user,SimpleUploadedFile('legacy.xlsx',content),'new',{'reason':'Legacy'})
        importing.commit(self.user,batch.pk)
        batch.mapping.pop('program_year');batch.save(update_fields=['mapping'])
        before=Demand.objects.count()
        repeated=importing.preview(self.user,SimpleUploadedFile('legacy.xlsx',content),'new',{'reason':'Legacy'})
        importing.commit(self.user,repeated.pk)
        self.assertEqual(repeated.pk,batch.pk);self.assertEqual(Demand.objects.count(),before)

    def test_revision_keeps_the_automatically_assigned_order_folio(self):
        batch=importing.preview(self.user,self.simple_template([[self.part.customer,'SLP',43,self.part.number,100]]),'new',{'reason':'New order'})
        importing.commit(self.user,batch.pk)
        batch.refresh_from_db();reference=batch.preview['rows'][0]['reference']
        workbook=Workbook();sheet=workbook.active;sheet.title='Programa'
        sheet.append(['cliente','destino_o_ambito','semana_programa','numero_parte','cantidad_programada','referencia'])
        sheet.append([self.part.customer,'SLP',43,self.part.number,250,reference])
        stream=BytesIO();workbook.save(stream)
        revision=importing.preview(self.user,SimpleUploadedFile('revision.xlsx',stream.getvalue()),'revision',{'reason':'Quantity revision'})
        self.assertFalse([i for i in revision.preview['issues'] if i['severity']=='blocking'])
        before=Demand.objects.count();importing.commit(self.user,revision.pk)
        line=Demand.objects.get(reference=reference)
        self.assertEqual(line.quantity,250);self.assertEqual(Demand.objects.count(),before)

    def test_manual_program_assigns_current_year_and_distinct_order_folios(self):
        response=self.client.post('/api/programs/',{'client':self.part.customer,'scope':'SLP','week':43,'year':1999,'reference':'IGNORED','group':self.group.pk,'lines':[{'part':self.part.pk,'quantity':100,'reference':'SAME'},{'part':self.part.pk,'quantity':200,'reference':'SAME'}]},format='json')
        self.assertEqual(response.status_code,201,response.data)
        program=Program.objects.get(pk=response.data['id'])
        self.assertEqual(program.year,timezone.localdate().year)
        self.assertEqual(program.group.name,self.part.customer)
        self.assertNotEqual(program.reference,'IGNORED')
        self.assertEqual(program.lines.values('reference').distinct().count(),2)
        for line in program.lines.all():self.assertRegex(line.reference,r'^ORD-\d{4}-\d{6,}$')

    def test_new_import_ignores_obsolete_group_year_and_reference(self):
        workbook=load_workbook(self.template_file());sheet=workbook.active
        sheet.cell(2,1,'OBSOLETE GROUP');sheet.cell(2,4,1999);sheet.cell(2,8,'OBSOLETE REF')
        stream=BytesIO();workbook.save(stream)
        batch=importing.preview(self.user,SimpleUploadedFile('old.xlsx',stream.getvalue()),'new',{'reason':'Old layout'})
        importing.commit(self.user,batch.pk)
        batch.refresh_from_db();row=batch.preview['rows'][0]
        self.assertEqual(row['year'],timezone.localdate().year)
        self.assertEqual(row['group'],self.part.customer)
        self.assertNotEqual(row['reference'],'OBSOLETE REF')

    def test_credit_compatible_across_legacy_groups_for_same_customer(self):
        event=self.pack(400);target=self.make_line(100,43)
        target.program.group=BalanceGroup.objects.create(name='OLD CLASSIFICATION')
        target.program.save(update_fields=['group'])
        self.assertEqual(s.transfer(self.user,event.pk,target.pk,100,'Customer only'),100)
        self.assertEqual(s.applied(target),100)

    def test_new_program_auto_selects_unique_template(self):
        batch=importing.preview(self.user,self.template_file(),'new',{'sheet':'','reason':'Prueba de detección automática'})
        self.assertEqual(batch.mapping['sheet'],'Programa')
        self.assertEqual(batch.preview['format'],'new_program')
        self.assertFalse(any(i['severity']=='blocking' for i in batch.preview['issues']))
        importing.commit(self.user,batch.pk)
        batch.refresh_from_db();self.assertTrue(batch.applied)

    def test_import_and_publish_is_visible_to_operator(self):
        batch=importing.preview(self.user,self.template_file(),'new',{'reason':'Prueba publicar'})
        importing.commit(self.user,batch.pk,publish=True)
        batch.refresh_from_db();result=importing.result(batch)
        self.assertTrue(result['visible_in_capture'])
        operator=get_user_model().objects.create_user(username='floor')
        operator.user_permissions.set(Permission.objects.filter(content_type__app_label='packing',codename__in=['consult','capture']))
        ids=[p['id'] for p in result['programs']]
        self.assertEqual(s.demand_queryset(operator,{}).filter(program_id__in=ids).count(),1)
        count=Demand.objects.count()
        importing.commit(self.user,batch.pk,publish=True)
        self.assertEqual(Demand.objects.count(),count)

    def test_saved_draft_can_be_published_from_same_batch(self):
        batch=importing.preview(self.user,self.template_file(),'new',{'reason':'Prueba borrador'})
        importing.commit(self.user,batch.pk)
        batch.refresh_from_db()
        self.assertFalse(importing.result(batch)['visible_in_capture'])
        importing.commit(self.user,batch.pk,publish=True)
        self.assertTrue(importing.result(batch)['visible_in_capture'])

    def test_inactive_group_publication_rolls_back_import(self):
        BalanceGroup.objects.create(name=self.part.customer,active=False)
        batch=importing.preview(self.user,self.template_file(),'new',{'reason':'Grupo inactivo'})
        count=Demand.objects.count()
        with self.assertRaises(ValidationError):importing.commit(self.user,batch.pk,publish=True)
        self.assertEqual(Demand.objects.count(),count)
        batch.refresh_from_db();self.assertFalse(batch.applied)

    def test_employee_setup_permissions_and_idempotency(self):
        response=self.client.post('/api/employees/',{'payroll':'NEW-P','name':'Empacador sintético nuevo'},format='json')
        self.assertEqual(response.status_code,201,response.data)
        self.assertEqual(self.client.post('/api/employees/',{'payroll':'NEW-P','name':'Empacador sintético nuevo'},format='json').status_code,200)
        self.assertEqual(Employee.objects.filter(payroll='NEW-P').count(),1)
        self.assertEqual(self.client.post('/api/employees/',{'payroll':'NEW-P','name':'Otra persona'},format='json').status_code,400)
        reader=get_user_model().objects.create_user(username='employee-reader')
        self.client.force_authenticate(reader)
        self.assertEqual(self.client.post('/api/employees/',{'payroll':'BLOCKED','name':'Sin permiso'},format='json').status_code,403)

    def test_multiple_templates_require_sheet_selection(self):
        file=self.template_file()
        w=load_workbook(file);w.copy_worksheet(w.active).title='Otro programa'
        out=BytesIO();w.save(out)
        result=importing.inspect(out.getvalue(),{},'new')
        self.assertEqual(len(result['rows']),0)
        self.assertTrue(any(i['severity']=='blocking' for i in result['issues']))

    def test_legacy_single_sheet_does_not_auto_select(self):
        w=Workbook();ws=w.active;ws.append(['CLIENTE','NÚMERO DE PARTE',42])
        out=BytesIO();w.save(out)
        result=importing.inspect(out.getvalue(),{},'new')
        self.assertTrue(any(i['severity']=='blocking' for i in result['issues']))

    def test_09_reconciliation_no_new_events(self):
        s.revise(self.user,self.line.pk,3200,'Prueba');self.pack(2500)
        self.assertEqual(self.line.quantity,300) # objeto original; el servicio carga y bloquea otra instancia
        result=s.balance(self.user,{})
        self.assertEqual(result['rows'][0]['pending'],700)
        w=Workbook();ws=w.active;ws.title='SEMANA 42';ws.append(['CLIENTE','DESTINO','NÚMERO DE PARTE','BALANCE']);ws.append(['CLIENTE-SINTÉTICO','SLP',self.part.number,700]);out=BytesIO();w.save(out)
        b=importing.preview(self.user,SimpleUploadedFile('reconcile.xlsx',out.getvalue()),'reconcile',{'sheet':ws.title,'group':self.group.name,'year':2026,'week':42,'program_id':self.line.program_id,'line_resolutions':{'2':self.line.pk},'balance_column':'D','reason':'Mismo corte'})
        importing.commit(self.user,b.pk);b.refresh_from_db()
        self.assertEqual(b.preview['reconciliation'][0]['difference'],0);self.assertEqual(Capture.objects.count(),1)

    def test_10_source_versions_blocked(self):
        w=Workbook();ws=w.active;ws.title='SEMANA 42 (2)';ws.append(['CLIENTE','NÚMERO DE PARTE',42]);ws.append(['CLIENTE-SINTÉTICO',self.part.number,100]);out=BytesIO();w.save(out)
        preview=importing.inspect(out.getvalue(),{'sheet':ws.title,'group':'SINTÉTICO','year':2026,'week':42,'quantity_columns':['C']},'new')
        self.assertTrue(any('versión' in i['message'] for i in preview['issues']))

    def test_11_ambiguous_week(self):
        path=Path(__file__).resolve().parents[2]/'BALANCE DAIKIN SLP, EUA.xlsx'
        if not path.exists():self.skipTest('Fuente real ausente en esta distribución')
        data=importing.inspect(path.read_bytes(),{'sheet':'SEMANA 42','group':'Daikin','year':2026,'week':42,'quantity_columns':['I'],'approve_provisional':True},'new')
        self.assertTrue(any('Semana del encabezado' in i['message'] and i['severity']=='blocking' for i in data['issues']))

    def test_12_simultaneous_captures_postgresql(self):
        s.revise(self.user,self.line.pk,150,'Prueba concurrente')
        payloads=[self.payload(100),self.payload(100)]
        def job(payload):
            close_old_connections()
            try:return s.capture(get_user_model().objects.get(pk=self.user.pk),payload).pk
            finally:close_old_connections()
        with ThreadPoolExecutor(2) as executor: ids=list(executor.map(job,payloads))
        self.assertEqual(len(set(ids)),2);self.assertEqual(s.applied(self.line),150)
        self.assertEqual(sum(s.available(e) for e in Capture.objects.all()),50)

    def test_13_retry_and_conflict(self):
        p=self.payload(100);a=s.capture(self.user,p);b=s.capture(self.user,p)
        self.assertEqual(a.pk,b.pk)
        with self.assertRaises(ValidationError):s.capture(self.user,{**p,'quantity':101})
        self.pack(100);self.assertEqual(Capture.objects.count(),2)

    def test_14_correction_reverts_credit_in_future(self):
        s.revise(self.user,self.line.pk,600,'Prueba');event=self.pack(650)
        future=self.make_line(400,43);s.transfer(self.user,event.pk,future.pk,50,'Prueba')
        new=s.correct(self.user,event.pk,600,'Error de 50',uuid.uuid4())
        event.refresh_from_db();self.assertTrue(event.voided);self.assertEqual(s.applied(self.line),600);self.assertEqual(s.applied(future),0);self.assertEqual(s.available(new),0)

    @override_settings(RELEASE_MODE='confirmation')
    def test_15_release_once(self):
        event=self.pending_capture(100);self.assertEqual(s.applied(self.line),0)
        s.release(self.user,event.pk,'TEST-I');s.release(self.user,event.pk,'TEST-I')
        self.assertEqual(s.applied(self.line),100)

    def test_16_revision_returns_credit(self):
        event=self.pack(250);s.revise(self.user,self.line.pk,200,'Reducción')
        self.assertEqual(s.applied(self.line),200);self.assertEqual(s.available(event),50)
        self.assertEqual(Capture.objects.count(),1)

    def test_17_night_and_late_capture(self):
        p=self.payload(10);late=timezone.make_aware(datetime.combine(timezone.localdate()-timedelta(days=1),dt_time(10)),timezone.get_current_timezone());p['effective_at']=late.isoformat();p['operational_date']=str(timezone.localdate()-timedelta(days=2))
        event=s.capture(self.user,p)
        self.assertNotEqual(event.created_at.date(),event.operational_date)

    def test_18_real_irregular_workbooks(self):
        root=Path(__file__).resolve().parents[2]
        count=0
        for path in root.glob('BALANCE*.xlsx'):
            w=load_workbook(path,data_only=False)
            sheet=next(s for s in w.sheetnames if 'SEMANA' in s)
            data=importing.inspect(path.read_bytes(),{'sheet':sheet,'group':'SINTÉTICO','year':2026,'week':42,'approve_provisional':True},'history')
            self.assertIn('omitted',data);self.assertIn('issues',data);count+=1
            for row in data['rows']:
                self.assertIsInstance(row['client'],str)
                self.assertNotIn(importing.norm(row['part']),['LUNES','MIERCOLES'])
                self.assertNotIn(importing.norm(row['client']),['PROGRAMADO','PENDIENTE / BALANCE'])
        if not count:self.skipTest('Fuentes reales no incluidas')
        self.assertEqual(count,4)

    def test_19_permissions_and_csrf(self):
        operator=get_user_model().objects.create_user(username='operator')
        operator.user_permissions.set(Permission.objects.filter(content_type__app_label='packing',codename__in=['consult','capture']))
        with self.assertRaises(PermissionDenied):s.revise(operator,self.line.pk,900,'Intento directo')
        self.client.force_authenticate(operator)
        self.assertEqual(self.client.post('/api/programs/',{},format='json').status_code,403)
        read=get_user_model().objects.create_user(username='tv');read.user_permissions.add(Permission.objects.get(content_type__app_label='packing',codename='consult'))
        with self.assertRaises(PermissionDenied):s.capture(read,self.payload(100))
        csrf=APIClient(enforce_csrf_checks=True)
        self.assertEqual(csrf.post('/api/session/',{'username':'x','password':'x'}).status_code,403)

    def test_20_export_cut_snapshot_immutable(self):
        self.pack(100);snap=s.snapshot(self.user,{})
        self.pack(100);snap.refresh_from_db()
        self.assertEqual(snap.payload['totals']['applied'],100)
        response=self.client.get('/api/export/',{'snapshot':snap.pk})
        ws=load_workbook(BytesIO(response.content),data_only=True).active
        self.assertEqual(ws.cell(2,11).value,100)

    def test_21_weighted_progress(self):
        s.revise(self.user,self.line.pk,100,'Prueba');self.pack(100)
        self.make_line(900,part=Part.objects.create(number='OTHER'))
        self.assertEqual(s.balance(self.user,{})['totals']['progress'],10)

    def test_22_preview_concurrent_watermark(self):
        batch=importing.preview(self.user,self.template_file(),'new',{'sheet':'Programa','reason':'Prueba'})
        self.pack(20)
        with self.assertRaises(ValidationError):importing.commit(self.user,batch.pk)
        self.assertEqual(Capture.objects.count(),1)

    def test_closed_period_blocks_correction_and_capture(self):
        event=self.pack(100);s.snapshot(self.user,{},close=True,reason='Cierre sintético')
        with self.assertRaises(ValidationError):self.pack(100)
        with self.assertRaises(ValidationError):s.correct(self.user,event.pk,50,'Prueba',uuid.uuid4())

    def test_api_end_to_end_program_capture_export(self):
        created=self.client.post('/api/programs/',{'group':self.group.pk,'client':'CLIENTE-SINTÉTICO','scope':'SLP','year':2026,'week':43,'reference':'E2E','lines':[{'part':self.part.pk,'quantity':300}]},format='json')
        self.assertEqual(created.status_code,201,created.data)
        pk=created.data['id'];self.assertEqual(self.client.post(f'/api/programs/{pk}/publication/',{'published':True},format='json').status_code,200)
        line=Demand.objects.get(program_id=pk)
        result=self.client.post('/api/captures/',self.payload(200,line),format='json');self.assertEqual(result.status_code,201,result.data)
        self.assertEqual(self.client.get('/api/balance/',{'year':2026,'week':43}).data['totals']['pending'],100)
        self.assertEqual(self.client.get('/api/export/',{'year':2026,'week':43}).status_code,200)

    def test_concurrent_credit_consumption(self):
        event=self.pack(400);future=self.make_line(100,43)
        def job(_):
            close_old_connections()
            try:
                try:return s.transfer(get_user_model().objects.get(pk=self.user.pk),event.pk,future.pk,100,'Concurrente')
                except ValidationError:return 0
            finally:close_old_connections()
        with ThreadPoolExecutor(2) as executor:result=list(executor.map(job,range(2)))
        self.assertEqual(sum(result),100);self.assertEqual(s.applied(future),100)

    def test_invalid_fractional_and_zero_quantities(self):
        for value in [0,-1,1.5,'1.5',True]:
            with self.assertRaises((ValidationError,ValueError)):s.capture(self.user,self.payload(value))

    def test_opening_balance_no_physical_events(self):
        file=self.template_file(3200)
        w=load_workbook(file);ws=w.active;ws.cell(1,9,'BALANCE');ws.cell(2,9,700);out=BytesIO();w.save(out)
        batch=importing.preview(self.user,SimpleUploadedFile('opening.xlsx',out.getvalue()),'opening',{'sheet':'Programa','balance_column':'I','cutoff':timezone.now().isoformat(),'reason':'Corte sintético'})
        importing.commit(self.user,batch.pk)
        opening=OpeningBalance.objects.get()
        self.assertEqual(opening.applied,2500);self.assertEqual(Capture.objects.count(),0)

    def test_representative_balance_query_budget(self):
        lines=Demand.objects.bulk_create([Demand(program=self.line.program,part=self.part,reference='LOAD-'+str(i),original=100,quantity=100) for i in range(1000)])
        user=self.user; packer=Employee.objects.get(payroll='TEST-P');inspector=Employee.objects.get(payroll='TEST-I')
        events=Capture.objects.bulk_create([Capture(key=uuid.uuid4(),fingerprint='synthetic',target=line,quantity=50,effective_at=timezone.now(),operational_date=timezone.localdate(),shift=self.shift,packer=packer,inspector=inspector,actor=user,released=True,release_mode='declared') for line in lines])
        Allocation.objects.bulk_create([Allocation(capture=event,demand=line,quantity=50) for line,event in zip(lines,events)])
        t=time.perf_counter()
        with CaptureQueriesContext(connection) as queries: result=s.balance(user,{})
        elapsed=time.perf_counter()-t
        self.assertEqual(len(result['rows']),1001)
        self.assertEqual(result['totals']['applied'],50000)
        self.assertLessEqual(len(queries),12)
        path=Path(__file__).resolve().parents[2]/'docs'/'evidence'/'performance.json';path.parent.mkdir(exist_ok=True)
        path.write_text(json.dumps({'synthetic':True,'database':'PostgreSQL','rows':1001,'captures':1000,'allocations':1000,'seconds':round(elapsed,4),'queries':len(queries),'scope':'Single balance request, local machine; not a user-capacity guarantee.'},indent=2),encoding='utf-8')
