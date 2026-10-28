import hashlib
import json
import re
import unicodedata
import zipfile
from io import BytesIO
from datetime import date
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from django.db import transaction
from django.utils import timezone
from rest_framework.exceptions import ValidationError
from .models import *
from .services import require, lock_ledger, audit, revise, applied, allocate, check_open
from .part_validation import resolve_part, resolve_part_by_number, validate_part_customer, validate_program
from .client_groups import client_group

MODES=['new','revision','opening','reconcile','history']

def client_group_name(client):
    # Compatibility key for the legacy FK; customers remain the only grouping.
    return client if len(client)<=100 else 'CLI-'+hashlib.sha256(client.encode()).hexdigest()

def norm(value):
    return ''.join(c for c in unicodedata.normalize('NFKD',str(value or '')) if not unicodedata.combining(c)).upper().strip()

def scalar(value):
    return value.isoformat() if hasattr(value,'isoformat') else value

def operational_identity(raw,client):
    if raw is None or not isinstance(client,str) or not client.strip():return False
    if hasattr(raw,'isoformat') or isinstance(raw,bool):return False
    if isinstance(raw,float) and not raw.is_integer():return False
    days={'LUNES','MARTES','MIERCOLES','JUEVES','VIERNES','SABADO','DOMINGO'}
    if norm(raw) in days or norm(client) in days:return False
    if norm(raw) in {'NUMERO DE PARTE','N/P','TOTAL','NUMERO_PARTE'}:return False
    if re.match(r'^(PROGRAMADO|PENDIENTE|ENTREGADO|AVANCE|TOTAL|META)(\b|\s|/)',norm(client)):return False
    if re.fullmatch(r'[WS]\s*\d+',norm(client)):return False
    return True

def watermark(): return Audit.objects.order_by('-id').values_list('id',flat=True).first() or 0

def inspect(content, mapping, mode):
    if len(content)>15*1024*1024: raise ValidationError('Archivo mayor a 15 MB.')
    try:
        with zipfile.ZipFile(BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist())>100*1024*1024: raise ValidationError('Contenido descomprimido mayor a 100 MB.')
        wf=load_workbook(BytesIO(content),data_only=False,keep_links=False)
        wv=load_workbook(BytesIO(content),data_only=True,keep_links=False)
    except (zipfile.BadZipFile,KeyError,ValueError): raise ValidationError('El archivo no es un XLSX válido.')
    result={'sheets':wf.sheetnames,'rows':[],'issues':[],'omitted':[],'controls':[],'watermark':watermark(),'cutoff':timezone.now().isoformat(),'mapping':mapping}
    sheet=mapping.get('sheet')
    if not sheet:
        # Solo las plantillas explícitas de demanda nueva pueden seleccionarse solas.
        # Un balance histórico de una sola hoja sigue exigiendo selección oficial.
        candidates=[]
        required={'SEMANA_PROGRAMA','NUMERO_PARTE','CANTIDAD_PROGRAMADA'}
        if mode in ['new','revision']:
            for candidate in wf:
                for row in candidate.iter_rows(max_row=min(30,candidate.max_row)):
                    if required.issubset({norm(c.value) for c in row if c.value is not None}):
                        candidates.append(candidate.title);break
        if len(candidates)==1:
            sheet=candidates[0]
            mapping={**mapping,'sheet':sheet}
            result['mapping']=mapping
            result['auto_selected_sheet']=True
        else:
            result['issues'].append({'severity':'blocking','message':'El archivo contiene varias plantillas. Selecciona la hoja que deseas importar.' if candidates else 'Selecciona una hoja y una versión oficial.'}); return result
    if sheet not in wf.sheetnames: raise ValidationError('Hoja inexistente.')
    sf=wf[sheet]; sv=wv[sheet]
    def issue(message,cell='',severity='blocking'):
        result['issues'].append({'sheet':sheet,'cell':cell,'severity':severity,'message':message})
    template=False; header=None; columns={}
    for row in sf.iter_rows(max_row=min(30,sf.max_row)):
        labels={norm(c.value):c.column for c in row if c.value is not None}
        if 'NUMERO_PARTE' in labels and 'CANTIDAD_PROGRAMADA' in labels:
            header=row[0].row; columns=labels; template=True; break
        if 'CLIENTE' in labels and any('PARTE' in l for l in labels):
            header=row[0].row; columns=labels; break
    if not header:
        issue('No se detectó encabezado de programa; los formularios no son demanda.'); return result
    result['format']='new_program' if template else 'legacy_balance'
    if not template and not re.search('SEMANA',norm(sheet)) and mode!='history': issue('La hoja no parece un programa semanal. Mantén como histórico o usa plantilla nueva.')
    def col(key,default=None):
        v=mapping.get(key,default)
        return column_index_from_string(v) if isinstance(v,str) and v.isalpha() else int(v) if v else None
    part_col=columns.get('NUMERO_PARTE') if template else next(v for k,v in columns.items() if 'PARTE' in k)
    client_col=columns.get('CLIENTE')
    scope_col=columns.get('DESTINO_O_AMBITO') if template else columns.get('DESTINO')
    quantity_cols=[columns['CANTIDAD_PROGRAMADA']] if template else [col('unused',v) for v in mapping.get('quantity_columns',[])]
    balance_col=col('balance_column',columns.get('BALANCE'))
    if not quantity_cols and mode not in ['history','reconcile']: issue('Selecciona columnas de demanda explícitas; no se suman arrastres automáticamente.')
    if mode in ['opening','reconcile'] and not balance_col: issue('Selecciona la columna de saldo de origen.')
    if not template:
        week_candidates=[(c.coordinate,c.value) for c in sf[header] if isinstance(c.value,(int,float)) and 1<=c.value<=53]
        title_weeks=re.findall(r'\d+',sheet)
        if title_weeks and week_candidates and int(title_weeks[0])!=int(week_candidates[0][1]):
            issue('Semana del encabezado distinta del nombre: requiere resolución explícita.',week_candidates[0][0],'warning' if mapping.get('resolve_week') else 'blocking')
        if ('(' in sheet or 'OK' in norm(sheet) or len(title_weeks)>1) and not mapping.get('resolve_version'): issue('Hoja con variante o semanas múltiples: selecciona versión oficial explícitamente.')
    seen=set(); seen_parts={}
    for r in range(header+1,sf.max_row+1):
        raw=sv.cell(r,part_col).value
        client=sv.cell(r,client_col).value if client_col else mapping.get('client')
        number=str(raw or '').strip()
        if isinstance(raw,(int,float)) and sf.cell(r,part_col).number_format.startswith('000'):
            number=f'{int(raw):0{len(sf.cell(r,part_col).number_format)}d}'
        selected=None
        catalog_identity=template and not str(client or '').strip()
        if catalog_identity and operational_identity(raw,'CATALOGO'):
            try:
                selected=resolve_part_by_number(number,mapping.get('part_resolutions',{}).get(str(r)))
                client=selected.customer
            except ValidationError as exc:
                issue(' '.join(str(item) for item in exc.detail),str(r))
                continue
        if not operational_identity(raw,'CATALOGO' if catalog_identity else client):
            if isinstance(client,str) and norm(client).startswith(('PROGRAMADO','PENDIENTE')) and client_col:
                control=sv.cell(r,client_col+1)
                if isinstance(control.value,(int,float)):
                    result['controls'].append({'label':client,'cell':control.coordinate,'value':control.value})
            result['omitted'].append({'row':r,'reason':'Sin identidad operativa válida: blanco, encabezado, resumen, calendario o nota auxiliar.','part_value':scalar(raw),'client_value':scalar(client)}); continue
        client=client_group(client)
        if sf.cell(r,part_col).data_type=='f': issue('La identidad de pieza no puede depender de una fórmula.',sf.cell(r,part_col).coordinate)
        scope=str(sv.cell(r,scope_col).value or '') if scope_col else str(mapping.get('scope',''))
        year=sv.cell(r,columns['ANIO_PROGRAMA']).value if template and 'ANIO_PROGRAMA' in columns else mapping.get('year')
        week=sv.cell(r,columns['SEMANA_PROGRAMA']).value if template and 'SEMANA_PROGRAMA' in columns else mapping.get('week')
        group=sv.cell(r,columns['GRUPO']).value if template and 'GRUPO' in columns else mapping.get('group')
        reference=str(sv.cell(r,columns['REFERENCIA']).value or number) if 'REFERENCIA' in columns else number
        if template and mode!='new':
            if not group: group=client_group_name(str(client).strip())
            if not year: year=timezone.localdate().year
            if mode=='revision' and 'REFERENCIA' not in columns:
                issue('Para revisar cantidades conserva el folio de la orden en la columna referencia.',str(r))
        if mode=='new':
            # New work is organized only by customer. Old columns are ignored.
            client=str(client).strip()
            group=client_group_name(client)
            year=timezone.localdate().year
            reference='ORD-'+hashlib.sha256(content).hexdigest()[:12].upper()+f'-{r-header:04d}'
        identity=(str(client),scope,str(year),str(week),reference)
        if identity in seen: issue('Renglón duplicado: proporciona referencia de pedido/renglón distinta.',sf.cell(r,part_col).coordinate)
        seen.add(identity)
        if not group or not year or not week: issue('Falta grupo, año o semana explícitos.',str(r))
        parsed_year=parsed_week=0
        try:
            parsed_year,parsed_week=int(year),int(week)
            if isinstance(week,bool) or float(week)!=parsed_week: raise ValueError
            date.fromisocalendar(parsed_year,parsed_week,1)
        except (ValueError,TypeError,OverflowError): issue('Año/semana no válidos.',str(r))
        evidence=[]; quantity=0; signed_balance=None
        relevant=set(quantity_cols+[part_col]+([balance_col] if balance_col else []))
        for c in sf[r]:
            v=sv[c.coordinate]
            if c.value is not None:
                evidence.append({'cell':c.coordinate,'original':scalar(c.value),'cached':scalar(v.value),'format':c.number_format})
            if v.data_type=='e': issue(str(v.value),c.coordinate,'blocking' if c.column in relevant else 'warning')
            if c.data_type=='f' and '[' in c.value: issue('Referencia externa conservada como evidencia; no ejecutada.',c.coordinate,'warning')
        for column in quantity_cols:
            v=sv.cell(r,column).value
            if v is None and sf.cell(r,column).data_type=='f': issue('Fórmula sin resultado almacenado en cantidad.',sf.cell(r,column).coordinate)
            elif v is None and template: issue('La cantidad programada es obligatoria.',sf.cell(r,column).coordinate)
            if v is not None:
                if isinstance(v,bool) or not isinstance(v,(int,float)) or int(v)!=v or v<0: issue('Cantidad inválida; no se sustituye por cero.',sv.cell(r,column).coordinate)
                else: quantity+=int(v)
        if balance_col:
            v=sv.cell(r,balance_col).value
            if isinstance(v,(int,float)) and not isinstance(v,bool) and int(v)==v: signed_balance=int(v)
            elif mode in ['opening','reconcile']: issue('Saldo de origen no es un entero almacenado.',sv.cell(r,balance_col).coordinate)
        resolutions=mapping.get('part_resolutions',{})
        try:
            selected=selected or resolve_part(number,client,resolutions.get(str(r)))
        except ValidationError as exc:
            if mode!='history': issue(' '.join(str(item) for item in exc.detail),str(r))
        if selected and mode=='new':
            client=client_group(selected.customer)
            number=selected.number
            group=client_group_name(client)
        if selected and template and not scope.strip():
            scope=selected.scope
        if mode=='opening' and signed_balance is not None and signed_balance>quantity: issue('Pendiente mayor que demanda mapeada; revisa arrastres/alcance.',str(r))
        record={'row':r,'group':str(group or ''),'client':str(client),'scope':scope,'year':parsed_year,'week':parsed_week,'part':number,'part_id':selected.pk if selected else None,'reference':reference,'quantity':quantity,'signed_balance':signed_balance,'evidence':evidence}
        p=Program.objects.filter(pk=mapping['program_id']).first() if mapping.get('program_id') else Program.objects.filter(reference=program_ref(record)).first()
        if mode=='new' and selected:
            part_key=(str(group),str(client),scope,parsed_year,parsed_week,selected.pk)
            if part_key in seen_parts:
                record['duplicate_in_file']=True
                record['duplicate_source_row']=seen_parts[part_key]
            else: seen_parts[part_key]=r
            matches=list(Demand.objects.filter(program=p,part=selected).order_by('id')) if p else []
            if len(matches)>1:
                issue('Ya existen varias órdenes de este número en el programa semanal; consolídalas antes de importar.',str(r))
            elif matches:
                found=matches[0]
                record.update(duplicate_existing=True,existing_demand=found.pk,current_quantity=found.quantity,quantity_after=found.quantity+quantity if mapping.get('duplicate_action','sum')=='sum' else quantity)
            if record.get('duplicate_in_file') or record.get('duplicate_existing'):
                action=mapping.get('duplicate_action','sum')
                if action not in ['sum','replace']:
                    issue('Elige si deseas sumar o sustituir la cantidad del número que ya existe.',str(r))
                else: record['duplicate_action']=action
        resolved=mapping.get('line_resolutions',{}).get(str(r))
        existing=(Demand.objects.filter(program=p,pk=resolved).first() if resolved else Demand.objects.filter(program=p,reference=reference).first()) if p else None
        if existing:
            from .services import available
            credit=sum(available(c) for c in Capture.objects.filter(target=existing,released=True,voided=False))+sum(available(o) for o in OpeningBalance.objects.filter(demand=existing))
            current=existing.quantity-applied(existing)-credit
            if mapping.get('snapshot'):
                snap=Snapshot.objects.filter(pk=mapping['snapshot']).first()
                match=next((v for v in snap.payload['rows'] if v['id']==existing.pk),None) if snap else None
                if match: current=match['quantity']-match['applied']-match['credit']
                else: issue('El corte no contiene el renglón.',str(r))
            record.update(existing_demand=existing.pk,current_quantity=existing.quantity,current_signed_balance=current,quantity_difference=quantity-existing.quantity,balance_difference=signed_balance-current if signed_balance is not None else None)
            if mode=='new': issue('Demanda existente: selecciona revisión en lugar de nuevo programa.',str(r))
        elif mode in ['revision','reconcile']: issue('No se encuentra el renglón; vincula program_id y line_resolutions.',str(r))
        result['rows'].append(record)
    result['totals']={'quantity':sum(r['quantity'] for r in result['rows']),'positive_pending':sum(max(r['signed_balance'] or 0,0) for r in result['rows']),'source_excess':sum(max(-(r['signed_balance'] or 0),0) for r in result['rows'])}
    for control in result['controls']:
        detail=result['totals']['quantity' if norm(control['label']).startswith('PROGRAMADO') else 'positive_pending']
        control['detail']=detail;control['difference']=control['value']-detail
        if control['difference']:
            issue(f'Resumen manual {control["label"]}: {control["value"]}; detalle mapeado: {detail}. Conciliar alcance antes de oficializar.',control['cell'],'warning')
    return result

@transaction.atomic
def preview(actor, file, mode, mapping):
    require(actor,'plan')
    lock_ledger()
    if mode not in MODES: raise ValidationError('Modo de importación inválido.')
    if not file.name.lower().endswith('.xlsx'): raise ValidationError('Solo archivos .xlsx; no se admiten macros.')
    content=file.read(); digest=hashlib.sha256(content).hexdigest()
    if mode=='new': mapping={**mapping,'program_year':timezone.localdate().year}
    data=inspect(content,mapping,mode)
    mapping=data['mapping']
    old=ImportBatch.objects.filter(digest=digest,mode=mode,mapping=mapping).first()
    if not old and mode=='new':
        legacy_mapping={key:value for key,value in mapping.items() if key!='program_year'}
        legacy=ImportBatch.objects.filter(digest=digest,mode=mode,mapping=legacy_mapping,applied=True).first()
        if legacy and all(row['year']==timezone.localdate().year for row in legacy.preview.get('rows',[])):
            old=legacy
    if old:
        if not old.applied:
            old.preview=data; old.save(update_fields=['preview'])
        return old
    file.seek(0)
    batch=ImportBatch.objects.create(digest=digest,file=file,mode=mode,mapping=mapping,preview=data,actor=actor)
    return batch

def program_ref(row):
    # Identidad operativa independiente de archivo/version: evita duplicar demanda revisada.
    return 'IMP-'+hashlib.sha256(json.dumps([row[k] for k in ['group','client','scope','year','week']],ensure_ascii=False).encode()).hexdigest()[:32]

@transaction.atomic
def commit(actor, pk, publish=False):
    require(actor,'plan'); lock_ledger(); check_open(timezone.localdate())
    batch=ImportBatch.objects.select_for_update().get(pk=pk)
    if batch.applied:
        if publish: publish_batch(actor,batch)
        return batch
    data=batch.preview
    if any(i['severity']=='blocking' for i in data['issues']): raise ValidationError('Resuelve los errores bloqueantes y genera otra previsualización.')
    if batch.mode=='new' and any(row['year']!=timezone.localdate().year for row in data['rows']):
        raise ValidationError('Cambió el año desde la revisión. Vuelve a revisar el archivo para asignar el año actual.')
    if watermark()!=data['watermark']: raise ValidationError('El estado cambió desde la previsualización; vuelve a calcularla antes de confirmar.')
    reason=batch.mapping.get('reason','')
    if not reason.strip(): raise ValidationError('Indica el motivo/decisión de Planeación.')
    differences=[]
    program_ids=set()
    if batch.mode=='history':
        batch.applied=True; batch.save(); audit(actor,'import_history',batch,reason); return batch
    for row in data['rows']:
        group,_=BalanceGroup.objects.get_or_create(name=row['group'])
        start=date.fromisocalendar(row['year'],row['week'],1)
        ref=program_ref(row)
        program=Program.objects.filter(pk=batch.mapping['program_id']).first() if batch.mapping.get('program_id') else Program.objects.filter(reference=ref).first()
        if program and (program.group_id!=group.pk or client_group(program.client)!=row['client'] or program.scope!=row['scope'] or program.year!=row['year'] or program.week!=row['week']):
            raise ValidationError('El programa vinculado no coincide con grupo, cliente, ámbito o semana del archivo.')
        if not program:
            if batch.mode in ['revision','reconcile']: raise ValidationError('No existe programa para revisar/conciliar; vincula el programa oficial.')
            program=Program.objects.create(group=group,client=row['client'],scope=row['scope'],year=row['year'],week=row['week'],starts_on=start,ends_on=date.fromisocalendar(row['year'],row['week'],7),reference=ref)
        program_ids.add(program.pk)
        if not row['part_id']: raise ValidationError('Cada pieza debe estar validada en Universo Ramos para el cliente de la orden.')
        part=Part.objects.get(pk=row['part_id'])
        validate_part_customer(part,row['client'])
        resolved=batch.mapping.get('line_resolutions',{}).get(str(row['row']))
        line=Demand.objects.filter(program=program,pk=resolved).first() if resolved else Demand.objects.filter(program=program,reference=row['reference']).first()
        if batch.mode=='new' and not line:
            matching=list(Demand.objects.select_for_update().filter(program=program,part=part).order_by('id'))
            if len(matching)>1: raise ValidationError('El programa tiene varias Ã³rdenes con este nÃºmero. ConsolÃ­dalas antes de importar.')
            line=matching[0] if matching else None
        existing_line=line is not None
        if batch.mode=='new' and line and line.part_id!=part.id: raise ValidationError('El número de parte existente no coincide con el archivo.')
        if batch.mode in ['revision','reconcile'] and not line: raise ValidationError('Renglón no encontrado en programa; revisa identidad y referencia.')
        if not line:
            line=Demand.objects.create(program=program,part=part,reference=row['reference'],original=row['quantity'],quantity=row['quantity'],attributes_snapshot={'number':part.number,'diameter':part.diameter,'attributes':part.attributes,'source':row['evidence']})
            if batch.mode=='new':
                line.reference=f'ORD-{program.year}-{line.pk:06d}'
                line.save(update_fields=['reference'])
                row['reference']=line.reference
        if line.part_id!=part.id: raise ValidationError('La pieza no coincide con la identidad del renglón existente.')
        if batch.mode=='new' and not resolved and existing_line:
            action=batch.mapping.get('duplicate_action','sum')
            if action not in ['sum','replace']: raise ValidationError('Elige si deseas sumar o sustituir la cantidad repetida y vuelve a revisar el Excel.')
            target=line.quantity+row['quantity'] if action=='sum' else row['quantity']
            revise(actor,line.pk,target,reason)
            line.original=line.original+row['quantity'] if action=='sum' else row['quantity']
            line.save(update_fields=['original'])
            row['reference']=line.reference
        elif batch.mode=='revision': revise(actor,line.pk,row['quantity'],reason)
        elif batch.mode=='opening':
            if Capture.objects.filter(target=line).exists() or line.allocations.exists() or OpeningBalance.objects.filter(demand=line).exists(): raise ValidationError('El saldo inicial exige un renglón sin operación ni saldo previo.')
            from .services import parse_time
            cutoff=parse_time(batch.mapping.get('cutoff',''))
            if cutoff>timezone.now(): raise ValidationError('El corte inicial no puede estar en el futuro.')
            source=OpeningBalance.objects.create(demand=line,applied=row['quantity']-max(row['signed_balance'],0),credit=max(-row['signed_balance'],0),cutoff=cutoff,evidence={'batch':batch.pk,'row':row['row']})
            if source.applied: Allocation.objects.create(opening=source,demand=line,quantity=source.applied)
        elif batch.mode=='reconcile':
            # Fotografía al momento de preview; watermark rechaza cualquier cambio concurrente.
            from .services import available
            credit=sum(available(c) for c in Capture.objects.filter(target=line,released=True,voided=False))+sum(available(o) for o in OpeningBalance.objects.filter(demand=line))
            official=line.quantity-applied(line)-credit
            if batch.mapping.get('snapshot'):
                snap=Snapshot.objects.get(pk=batch.mapping['snapshot'])
                match=next((r for r in snap.payload['rows'] if r['id']==line.id),None)
                if not match: raise ValidationError('El corte seleccionado no contiene el renglón.')
                official=match['quantity']-match['applied']-match['credit']
            differences.append({'demand':line.pk,'excel':row['signed_balance'],'calculated':official,'difference':row['signed_balance']-official,'explanation':reason})
    data['reconciliation']=differences
    data['program_ids']=sorted(program_ids)
    batch.preview=data; batch.applied=True; batch.save()
    audit(actor,'import_'+batch.mode,batch,reason,totals=data['totals'],differences=differences)
    if publish: publish_batch(actor,batch)
    return batch

def batch_programs(batch):
    if not batch.applied or batch.mode=='history':return Program.objects.none()
    ids=batch.preview.get('program_ids')
    if ids is not None:return Program.objects.filter(pk__in=ids).order_by('id')
    # Compatibilidad con lotes confirmados antes de guardar sus IDs de programa.
    if batch.mapping.get('program_id'):return Program.objects.filter(pk=batch.mapping['program_id'])
    return Program.objects.filter(reference__in={program_ref(row) for row in batch.preview.get('rows',[])}).order_by('id')

def publish_batch(actor,batch):
    require(actor,'plan')
    if batch.mode not in ['new','opening']:raise ValidationError('Esta carga no crea órdenes publicables.')
    if result(batch)['missing_programs']:
        raise ValidationError('Esta carga contiene programas eliminados. El lote conserva su evidencia y no puede volver a publicarlos.')
    programs=list(batch_programs(batch).select_related('group'))
    if not programs:raise ValidationError('No hay programas guardados para publicar.')
    for program in programs:
        validate_program(program)
        if not program.group.active:raise ValidationError('El grupo está desactivado para publicaciones.')
        if not program.published:
            program.published=True;program.save(update_fields=['published'])
            audit(actor,'publication',program,'Publicación desde importación revisada',batch=batch.pk,published=True)

def result(batch):
    programs=list(batch_programs(batch).values('id','client','scope','year','week','published','group_id','group__name'))
    missing=[]
    if batch.applied and batch.mode in ['new','opening']:
        expected=batch.preview.get('program_ids')
        if expected is not None:missing=sorted(set(expected)-{p['id'] for p in programs})
        elif batch.mapping.get('program_id') and not programs:missing=[batch.mapping['program_id']]
        elif not batch.mapping.get('program_id'):
            references={program_ref(row) for row in batch.preview.get('rows',[])}
            missing=sorted(references-set(Program.objects.filter(reference__in=references).values_list('reference',flat=True)))
    return {'id':batch.pk,'mode':batch.mode,'applied':batch.applied,'preview':batch.preview,'programs':programs,'missing_programs':missing,'visible_in_capture':bool(programs) and not missing and all(p['published'] for p in programs)}
