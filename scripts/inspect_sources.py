"""Inventario reproducible; no modifica fuentes ni recalcula fórmulas."""
import hashlib
import json
import re
import zipfile
from pathlib import Path
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
inventory=[]
for file in root.glob('*.xlsx'):
    formulas=load_workbook(file,data_only=False)
    values=load_workbook(file,data_only=True)
    book={'file':file.name,'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'sheets':[]}
    for sheet in formulas:
        errors=[]; external=[]; headers=[]; formula_count=0
        for row in sheet:
            for cell in row:
                value=values[sheet.title][cell.coordinate]
                if value.data_type=='e': errors.append({'cell':cell.coordinate,'value':value.value})
                if cell.data_type=='f':
                    formula_count+=1
                    if '[' in cell.value: external.append({'cell':cell.coordinate,'formula':cell.value})
            if row[0].row<=20:
                parts=[{'cell':c.coordinate,'value':str(c.value)} for c in row if c.value is not None]
                if parts: headers.append(parts)
        book['sheets'].append({'name':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'merged':[str(r) for r in sheet.merged_cells.ranges],'headers':headers,'formulas':formula_count,'external':external,'errors':errors,'program_candidate':bool(re.search('SEMANA',sheet.title,re.I))})
    inventory.append(book)
pbix=[]
for file in root.glob('*.pbix'):
    with zipfile.ZipFile(file) as z:
        layout=json.loads(z.read('Report/Layout').decode('utf-16-le'))
        pbix.append({'file':file.name,'entries':z.namelist(),'layout':layout})
(root/'docs'/'inventario_fuentes.json').write_text(json.dumps({'workbooks':inventory,'pbix':pbix},ensure_ascii=False,indent=2),encoding='utf-8')
print('Inventario:',len(inventory),'libros;',sum(len(b['sheets']) for b in inventory),'hojas;',sum(len(s['errors']) for b in inventory for s in b['sheets']),'errores almacenados.')
